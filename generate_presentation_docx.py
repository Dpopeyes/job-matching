import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_presentation_doc(output_path):
    doc = docx.Document()

    # Set standard margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("ENGLISH PRESENTATION SCRIPT\nBLUEHOUSE JOBS PLATFORM")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p_title.paragraph_format.space_after = Pt(4)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Project Overview: Sections 1 & 2 (Introduction, Objectives & Tools Used)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    p_sub.paragraph_format.space_after = Pt(20)

    # Metadata Card Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        [("Project Name:", True), ("BlueHouse Jobs (Smart Job Matching Platform)", False)],
        [("Presentation Focus:", True), ("Section 1: Origin & Objectives | Section 2: Technologies & Tools", False)]
    ]

    for r_idx, row_content in enumerate(meta_data):
        for c_idx, (text, is_bold) in enumerate(row_content):
            cell = meta_table.cell(r_idx, c_idx)
            cell.width = Inches(4.2) if c_idx == 1 else Inches(2.2)
            set_cell_background(cell, "F1F5F9" if r_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run.font.bold = is_bold
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Presentation Slides Content Structure
    slides = [
        {
            "slide_no": "Slide 1",
            "section": "Greeting & Introduction (การทักทายและแนะนำตัว)",
            "topic": "Welcome & Project Overview",
            "script": "Good morning/afternoon, respected committee members and audience. Welcome to our presentation.\n\nToday, we are excited to introduce our project called 'BlueHouse Jobs'—a Smart Job Matching Platform designed specifically for fresh graduates and employers.",
            "thai_notes": "สวัสดีครับ/ค่ะ ท่านกรรมการและผู้ฟังทุกท่าน ขอต้อนรับสู่การนำเสนอโครงงาน วันนี้พวกเรามีความยินดีอย่างยิ่งที่จะมานำเสนอโครงการ 'BlueHouse Jobs' แพลตฟอร์มคัดกรองจัดหางานอัจฉริยะสำหรับนักศึกษาจบใหม่และองค์กรนายจ้าง",
            "speaking_tips": "พูดด้วยน้ำเสียงแจ่มใส มั่นใจ สบตากรรมการ และเน้นเสียงคำว่า 'BlueHouse Jobs' และ 'Smart Job Matching Platform'"
        },
        {
            "slide_no": "Slide 2",
            "section": "Section 1: Project Origin & Problem Statement (ที่มาและความสำคัญ)",
            "topic": "Why BlueHouse Jobs?",
            "script": "Let us start with Section 1: The background and origin of our project.\n\nIn today's digital era, finding a job is a critical milestone for fresh graduates. However, many job seekers face two major problems: first, traditional platforms have cluttered, complicated interfaces; and second, there is a mismatch between students' actual skills and employer requirements.\n\nTo solve these challenges, we developed BlueHouse Jobs as a centralized, intuitive web platform that connects job seekers, employers, and administrators seamlessly.",
            "thai_notes": "ขอเริ่มต้นที่ข้อ 1 ที่มาและความสำคัญของโครงการ ในยุคดิจิทัล การหางานเป็นก้าวสำคัญของนักศึกษาจบใหม่ แต่หลายคนเจอปัญหาใหญ่ 2 ข้อ คือ เว็บเดิมๆ ใช้งานยากซับซ้อน และทักษะของผู้สมัครไม่ตรงกับที่นายจ้างต้องการ เราจึงพัฒนา BlueHouse Jobs ขึ้นมาเพื่อเป็นศูนย์กลางที่ใช้งานง่ายและเชื่อมโยงทุกคนเข้าด้วยกัน",
            "speaking_tips": "เน้นเสียงหนักแน่นตรงคำว่า 'two major problems', 'cluttered interfaces', และ 'mismatch' เพื่อชี้ให้เห็นปัญหาที่ระบบเราเข้ามาแก้ไข"
        },
        {
            "slide_no": "Slide 3",
            "section": "Section 1.1: Core Objectives (วัตถุประสงค์หลักของระบบ)",
            "topic": "Project Objectives",
            "script": "Based on these problems, our project has four primary objectives:\n\n1. First, to create a modern, responsive web application accessible from any device.\n2. Second, to help fresh graduates easily search for jobs based on location, salary, and categories, and submit applications effortlessly.\n3. Third, to enable employers to post job vacancies, manage listings, and review applicants systematically.\n4. And fourth, to provide an efficient Admin Moderation System to ensure platform safety and content quality.",
            "thai_notes": "จากปัญหาข้างต้น โครงการเรามีวัตถุประสงค์หลัก 4 ข้อ ได้แก่: 1) สร้างเว็บแอปพลิเคชันที่ทันสมัยรองรับทุกอุปกรณ์ 2) ช่วยให้นักศึกษาค้นหางานตามเงื่อนไขและยื่นใบสมัครได้สะดวก 3) ช่วยให้นายจ้างลงประกาศและคัดเลือกผู้สมัครได้อย่างเป็นระบบ 4) ให้บริการระบบผู้ดูแลระบบ (Admin) เพื่อกำกับดูแลความปลอดภัยของแพลตฟอร์ม",
            "speaking_tips": "ใช้นิ้วมือช่วยนับตามลำดับ 1, 2, 3, 4 ช่วยให้ฟังง่ายและกรรมการติดตามเนื้อหาได้ทัน"
        },
        {
            "slide_no": "Slide 4",
            "section": "Section 2: Technologies & Tools Used (เครื่องมือและเทคโนโลยีในการพัฒนา)",
            "topic": "Frontend & Backend Architecture",
            "script": "Moving on to Section 2: Technologies and tools used in development.\n\nBlueHouse Jobs is built using a modern Full-Stack Web Development architecture:\n\n- On the Frontend, we use React.js powered by Vite as a Single Page Application, providing lightning-fast rendering and smooth dynamic navigation without page reloads.\n- For Styling, we implemented a custom Vanilla CSS3 design system with dynamic gradients, combined with Lucide-React icons for a clean, modern user experience.\n- On the Backend, we built a Node.js and Express.js RESTful API to handle business logic, role-based access control, and server authentication.",
            "thai_notes": "มาต่อกันที่ข้อ 2 เทคโนโลยีและเครื่องมือที่ใช้พัฒนา ระบบเราพัฒนาด้วยสถาปัตยกรรม Full-Stack Web: ฝั่ง Frontend ใช้ React.js ร่วมกับ Vite ให้ความเร็วสูง โหลดหน้าเว็บแบบ SPA ไม่ต้องรีเฟรช / ส่วนดีไซน์ใช้ Vanilla CSS3 แบบ Custom พร้อมไอคอน Lucide-React / ฝั่ง Backend ใช้ Node.js และ Express.js ทำหน้าที่เป็น RESTful API จัดการสิทธิ์และตรรกะการทำงาน",
            "speaking_tips": "ออกเสียงคำศัพท์เทคนิคให้ชัดเจน: React.js (รีแอค-เจเอส), Vite (วีท), Single Page Application (ซิงเกิล-เพจ-แอปพลิเคชัน), RESTful API (เรสต์ฟูล-เอพีไอ)"
        },
        {
            "slide_no": "Slide 5",
            "section": "Section 2: Database & AI Engine (ฐานข้อมูลและระบบ AI Matching)",
            "topic": "Database & Smart AI Matching",
            "script": "For data management and intelligent features:\n\n- We use SQLite3 as our universal relational database to safely store user profiles, job posts, applications, and chat messages.\n- Furthermore, we developed a Dynamic AI Matching Engine that calculates a personalized Match Rate percentage based on applicant majors and specialized skills.\n- We also integrated the Google Gemini AI API to provide real-time, in-depth Thai language candidate analysis.\n- Finally, Git and GitHub Cloud were used for version control and source code collaboration.",
            "thai_notes": "ในส่วนการจัดการข้อมูลและฟีเจอร์อัจฉริยะ: เราใช้ SQLite3 เป็นฐานข้อมูลจัดเก็บข้อมูลผู้ใช้ โพสต์งาน ใบสมัคร และแชท / นอกจากนี้ เราได้พัฒนาระบบ AI Matching คำนวณเปอร์เซ็นต์ Match Rate % ตามสาขาวิชาและทักษะจริง / พร้อมเชื่อมต่อ Google Gemini AI API เพื่อวิเคราะห์ผู้สมัครเป็นภาษาไทยแบบเรียลไทม์ / และใช้ Git/GitHub ควบคุมเวอร์ชันซอร์สโค้ด",
            "speaking_tips": "เน้นฟีเจอร์ไฮไลท์ตรง 'Dynamic AI Matching Engine' และ 'Google Gemini AI API' เพราะเป็นจุดเด่นสำคัญของงานเรา"
        },
        {
            "slide_no": "Slide 6",
            "section": "Conclusion & Demo Transition (สรุปและเกริ่นเข้าสู่การสาธิตระบบ)",
            "topic": "Summary & Live Demonstration",
            "script": "In summary, BlueHouse Jobs effectively addresses the gap between fresh graduates and employers through a modern, secure, and AI-powered platform.\n\nNow, we would like to invite you to watch a live demonstration of our application in action. Thank you very much, and we are ready for your questions.",
            "thai_notes": "โดยสรุปแล้ว BlueHouse Jobs ตอบโจทย์การเชื่อมโยงนักศึกษาจบใหม่กับนายจ้างด้วยแพลตฟอร์มที่ทันสมัย ปลอดภัย และมีระบบ AI ช่วยแมตช์งาน ลำดับต่อไป ขอเชิญรับชมการสาธิตการใช้งานระบบจริง (Live Demo) ขอบคุณครับ/ค่ะ และพร้อมรับคำถามจากคณะกรรมการครับ/ค่ะ",
            "speaking_tips": "กล่าวโค้งขอบคุณอย่างสุภาพ ผายมือไปที่หน้าจอสำหรับการสาธิตระบบจริง (Live Demo)"
        }
    ]

    for slide in slides:
        # Section Heading Box
        p_sec = doc.add_paragraph()
        run_sec = p_sec.add_run(f"P {slide['slide_no']}: {slide['section']}")
        run_sec.font.name = "Arial"
        run_sec.font.size = Pt(13)
        run_sec.font.bold = True
        run_sec.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        p_sec.paragraph_format.space_before = Pt(14)
        p_sec.paragraph_format.space_after = Pt(6)

        # Table for Script Details
        table = doc.add_table(rows=4, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        cell_data = [
            ("Slide Topic / หัวข้อ", slide['topic']),
            ("English Speech Script\n(บทพูดนำเสนอภาษาอังกฤษ)", slide['script']),
            ("Thai Translation\n(คำแปลและความหมายภาษาไทย)", slide['thai_notes']),
            ("Presentation & Vocal Tips\n(เทคนิคการพูดและนำเสนอ)", slide['speaking_tips'])
        ]

        for r_idx, (label, val) in enumerate(cell_data):
            cell_lbl = table.cell(r_idx, 0)
            cell_val = table.cell(r_idx, 1)

            cell_lbl.width = Inches(2.0)
            cell_val.width = Inches(4.4)

            # Background styling
            set_cell_background(cell_lbl, "1E3A8A" if r_idx == 0 else "F8FAFC")
            set_cell_background(cell_val, "EFF6FF" if r_idx == 1 else ("FFFFFF" if r_idx % 2 == 0 else "F8FAFC"))

            set_cell_margins(cell_lbl, top=120, bottom=120, left=140, right=140)
            set_cell_margins(cell_val, top=120, bottom=120, left=140, right=140)

            # Label text
            p_l = cell_lbl.paragraphs[0]
            p_l.paragraph_format.space_after = Pt(2)
            run_l = p_l.add_run(label)
            run_l.font.name = "Arial"
            run_l.font.size = Pt(9.5)
            run_l.font.bold = True
            run_l.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r_idx == 0 else RGBColor(0x1E, 0x3A, 0x8A)

            # Value text
            p_v = cell_val.paragraphs[0]
            p_v.paragraph_format.space_after = Pt(2)
            run_v = p_v.add_run(val)
            run_v.font.name = "Arial"
            run_v.font.size = Pt(10)
            if r_idx == 1:
                run_v.font.bold = True
                run_v.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif r_idx == 3:
                run_v.font.italic = True
                run_v.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
            else:
                run_v.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Useful Vocabulary & Phrases Section
    p_vocab_hdr = doc.add_paragraph()
    run_vh = p_vocab_hdr.add_run("KEY VOCABULARY & TECHNICAL TERMS (คำศัพท์สำคัญที่ใช้ในการนำเสนอ)")
    run_vh.font.name = "Arial"
    run_vh.font.size = Pt(13)
    run_vh.font.bold = True
    run_vh.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p_vocab_hdr.paragraph_format.space_before = Pt(16)
    p_vocab_hdr.paragraph_format.space_after = Pt(8)

    vocab_table = doc.add_table(rows=9, cols=3)
    vocab_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    vocab_table.autofit = False

    vocab_headers = ["English Term", "Thai Translation", "Pronunciation / Usage Tip"]
    for c_idx, h_text in enumerate(vocab_headers):
        cell = vocab_table.cell(0, c_idx)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        run = p.add_run(h_text)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    vocab_items = [
        ("Smart Job Matching Platform", "แพลตฟอร์มคัดกรองจัดหางานอัจฉริยะ", "สมาร์ต-จ็อบ-แมตชิง-แพลตฟอร์ม"),
        ("Fresh Graduates", "นักศึกษาจบใหม่ / ผู้สำเร็จการศึกษาใหม่", "เฟรช-แกรด-จู-เอทส์"),
        ("Mismatch", "ความไม่สอดคล้อง / ไม่ตรงกันของทักษะ", "มิส-แมตช์"),
        ("Single Page Application (SPA)", "เว็บแอปพลิเคชันหน้าเดียวที่โหลดรวดเร็ว", "ซิงเกิล-เพจ-แอปพลิเคชัน"),
        ("RESTful API", "สถาปัตยกรรมเชื่อมต่อข้อมูลระหว่างหน้าเว็บกับเซิร์ฟเวอร์", "เรสต์ฟูล-เอพีไอ"),
        ("Role-Based Access Control (RBAC)", "ระบบควบคุมสิทธิ์การใช้งานตามบทบาท (User/Employer/Admin)", "โรล-เบสด์-แอคเซส-คอนโทรล"),
        ("Dynamic AI Matching Engine", "ระบบเอไอคำนวณอัตราความเข้ากันได้แบบเรียลไทม์", "ไดนามิก-เอไอ-แมตชิง-เอนจิน"),
        ("Admin Moderation System", "ระบบผู้ดูแลระบบสำหรับกำกับดูแลความปลอดภัยของเนื้อหา", "แอดมิน-มอดเดอเรชัน-ซิสเต็ม")
    ]

    for r_idx, (term, th_trans, pron) in enumerate(vocab_items, start=1):
        row_cells = [vocab_table.cell(r_idx, 0), vocab_table.cell(r_idx, 1), vocab_table.cell(r_idx, 2)]
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        widths = [Inches(2.3), Inches(2.3), Inches(1.8)]

        for c_idx, cell in enumerate(row_cells):
            cell.width = widths[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            val = [term, th_trans, pron][c_idx]
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            if c_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            elif c_idx == 2:
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Save document
    doc.save(output_path)
    print("SUCCESS: Presentation docx generated successfully at " + output_path)

if __name__ == "__main__":
    create_presentation_doc("d:/work90/bluehouse/BlueHouse_Jobs_Presentation_Script.docx")
