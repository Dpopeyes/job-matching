import docx
import sys
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_enhanced_docx(output_path):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Document Header Title Card
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("ENGLISH PRESENTATION SCRIPT & SLIDE GUIDE\nBLUEHOUSE JOBS PLATFORM")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p_title.paragraph_format.space_after = Pt(4)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Project Overview: Section 1 (Origin & Objectives) & Section 2 (Tech Stack & AI Architecture)")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    p_sub.paragraph_format.space_after = Pt(16)

    # Metadata Table Box
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        [("Project Title:", "BlueHouse Jobs (Smart Job Matching Platform)"), ("Version & Stack:", "v2.5 Full-Stack Web + AI Engine")],
        [("Presentation Focus:", "Sections 1 & 2 of Project Report"), ("Target Audience:", "Committee & General Presentation")]
    ]

    for r_idx, row in enumerate(meta_items):
        for c_idx, (lbl, val) in enumerate(row):
            cell = meta_table.cell(r_idx, c_idx)
            cell.width = Inches(3.3)
            set_cell_background(cell, "F1F5F9" if r_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r1 = p.add_run(f"{lbl} ")
            r1.font.name = "Arial"
            r1.font.size = Pt(9.5)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

            r2 = p.add_run(val)
            r2.font.name = "Arial"
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Detailed Slide-by-Slide Sections
    slides = [
        {
            "slide_no": "Slide 1",
            "section_title": "Greeting & Project Opening (เกริ่นนำและแนะนำตัว)",
            "topic": "Welcome & Project Introduction",
            "speech": "Good morning/afternoon, respected committee members and audience. Welcome to our presentation.\n\nToday, we are excited to introduce our project called 'BlueHouse Jobs'—a Smart Job Matching Platform designed specifically for fresh graduates and employers.",
            "translation": "สวัสดีครับ/ค่ะ ท่านกรรมการและผู้ฟังทุกท่าน ขอต้อนรับสู่การนำเสนอโครงงาน วันนี้พวกเรามีความยินดีอย่างยิ่งที่จะมานำเสนอโครงการ 'BlueHouse Jobs' แพลตฟอร์มคัดกรองจัดหางานอัจฉริยะสำหรับนักศึกษาจบใหม่และองค์กรนายจ้าง",
            "tips": "พูดด้วยน้ำเสียงสดใส ชัดเจน และสบตากรรมการ เน้นเสียงหนักแน่นตรงชื่อโครงการ 'BlueHouse Jobs'"
        },
        {
            "slide_no": "Slide 2",
            "section_title": "Section 1: Background & Problem Statement (ที่มาและความสำคัญ)",
            "topic": "Why BlueHouse Jobs?",
            "speech": "Let us start with Section 1: The background and origin of our project.\n\nIn today's digital era, finding a job is a critical milestone for fresh graduates. However, many job seekers face two major problems: first, traditional platforms have cluttered, complicated interfaces; and second, there is a mismatch between students' actual skills and employer requirements.\n\nTo solve these challenges, we developed BlueHouse Jobs as a centralized, intuitive web platform that connects job seekers, employers, and administrators seamlessly.",
            "translation": "ขอเริ่มต้นที่ข้อ 1 ที่มาและความสำคัญของโครงการ ในยุคดิจิทัล การหางานเป็นก้าวสำคัญของนักศึกษาจบใหม่ แต่หลายคนเจอปัญหาใหญ่ 2 ข้อ คือ เว็บเดิมๆ ใช้งานยากซับซ้อน และทักษะของผู้สมัครไม่ตรงกับที่นายจ้างต้องการ เราจึงพัฒนา BlueHouse Jobs ขึ้นมาเพื่อเป็นศูนย์กลางที่ใช้งานง่ายและเชื่อมโยงทุกคนเข้าด้วยกัน",
            "tips": "เน้นเสียงหนักตรงคำว่า 'two major problems', 'cluttered interfaces', และ 'mismatch' เพื่อเน้นย้ำถึงปัญหา"
        },
        {
            "slide_no": "Slide 3",
            "section_title": "Section 1.1: Core Objectives (วัตถุประสงค์หลักของระบบ)",
            "topic": "Four Primary Project Objectives",
            "speech": "Based on these problems, our project has four primary objectives:\n\n1. First, to create a modern, responsive web application accessible from any device.\n2. Second, to help fresh graduates easily search for jobs based on location, salary, and categories, and submit applications effortlessly.\n3. Third, to enable employers to post job vacancies, manage listings, and review applicants systematically.\n4. And fourth, to provide an efficient Admin Moderation System to ensure platform safety and content quality.",
            "translation": "จากปัญหาข้างต้น โครงการเรามีวัตถุประสงค์หลัก 4 ข้อ ได้แก่: 1) สร้างเว็บแอปพลิเคชันที่ทันสมัยรองรับทุกอุปกรณ์ 2) ช่วยให้นักศึกษาค้นหางานตามเงื่อนไขและยื่นใบสมัครได้สะดวก 3) ช่วยให้นายจ้างลงประกาศและคัดเลือกผู้สมัครได้อย่างเป็นระบบ 4) ให้บริการระบบผู้ดูแลระบบ (Admin) เพื่อกำกับดูแลความปลอดภัยของแพลตฟอร์ม",
            "tips": "ใช้นิ้วมือช่วยนับตามลำดับ 1, 2, 3, 4 ช่วยให้ฟังง่ายและกรรมการติดตามเนื้อหาได้ทัน"
        },
        {
            "slide_no": "Slide 4",
            "section_title": "Section 2: Technologies & Tools (เครื่องมือและเทคโนโลยีที่ใช้พัฒนา)",
            "topic": "Frontend & Backend Web Architecture",
            "speech": "Moving on to Section 2: Technologies and tools used in development.\n\nBlueHouse Jobs is built using a modern Full-Stack Web Development architecture:\n\n- On the Frontend, we use React.js powered by Vite as a Single Page Application, providing lightning-fast rendering and smooth dynamic navigation without page reloads.\n- For Styling, we implemented a custom Vanilla CSS3 design system with dynamic gradients, combined with Lucide-React icons for a clean, modern user experience.\n- On the Backend, we built a Node.js and Express.js RESTful API to handle business logic, role-based access control, and server authentication.",
            "translation": "มาต่อกันที่ข้อ 2 เทคโนโลยีและเครื่องมือที่ใช้พัฒนา ระบบเราพัฒนาด้วยสถาปัตยกรรม Full-Stack Web: ฝั่ง Frontend ใช้ React.js ร่วมกับ Vite ให้ความเร็วสูง โหลดหน้าเว็บแบบ SPA ไม่ต้องรีเฟรช / ส่วนดีไซน์ใช้ Vanilla CSS3 แบบ Custom พร้อมไอคอน Lucide-React / ฝั่ง Backend ใช้ Node.js และ Express.js ทำหน้าที่เป็น RESTful API จัดการสิทธิ์และตรรกะการทำงาน",
            "tips": "ออกเสียงคำเทคนิคชัดๆ: React.js (รีแอค-เจเอส), Vite (วีท), Single Page Application (ซิงเกิล-เพจ-แอป), RESTful API (เรสต์ฟูล-เอพีไอ)"
        },
        {
            "slide_no": "Slide 5",
            "section_title": "Section 2: Database & AI Engine (ฐานข้อมูล SQLite3 & ระบบ AI Matching)",
            "topic": "Intelligent Candidate Screening & Gemini AI Integration",
            "speech": "For data management and intelligent features:\n\n- We use SQLite3 as our universal relational database to safely store user profiles, job posts, applications, and chat messages.\n- Furthermore, we developed a Dynamic AI Matching Engine that calculates a personalized Match Rate percentage based on applicant majors and specialized skills.\n- We also integrated the Google Gemini AI API to provide real-time, in-depth Thai language candidate analysis.\n- Finally, Git and GitHub Cloud were used for version control and source code collaboration.",
            "translation": "ในส่วนการจัดการข้อมูลและฟีเจอร์อัจฉริยะ: เราใช้ SQLite3 เป็นฐานข้อมูลจัดเก็บข้อมูลผู้ใช้ โพสต์งาน ใบสมัคร และแชท / นอกจากนี้ เราได้พัฒนาระบบ AI Matching คำนวณเปอร์เซ็นต์ Match Rate % ตามสาขาวิชาและทักษะจริง / พร้อมเชื่อมต่อ Google Gemini AI API เพื่อวิเคราะห์ผู้สมัครเป็นภาษาไทยแบบเรียลไทม์ / และใช้ Git/GitHub ควบคุมเวอร์ชันซอร์สโค้ด",
            "tips": "เน้นฟีเจอร์ไฮไลท์ตรง 'Dynamic AI Matching Engine' และ 'Google Gemini AI API' เพราะเป็นจุดขายสำคัญของงาน"
        },
        {
            "slide_no": "Slide 6",
            "section_title": "Conclusion & Presentation Closing (สรุปและจบการนำเสนอ)",
            "topic": "Summary & Presentation Conclusion",
            "speech": "In summary, BlueHouse Jobs effectively addresses the gap between fresh graduates and employers through a modern, secure, and AI-powered platform.\n\nThat concludes our presentation for today. Thank you very much for your time and attention, and we are now ready for your questions.",
            "translation": "โดยสรุปแล้ว BlueHouse Jobs ตอบโจทย์การเชื่อมโยงนักศึกษาจบใหม่กับนายจ้างด้วยแพลตฟอร์มที่ทันสมัย ปลอดภัย และมีระบบ AI ช่วยแมตช์งาน พวกเราขอจบการนำเสนอเพียงเท่านี้ ขอบคุณครับ/ค่ะ สำหรับการรับฟัง และพร้อมรับคำถามจากคณะกรรมการครับ/ค่ะ",
            "tips": "กล่าวโค้งขอบคุณอย่างสุภาพและสบตากรรมการ"
        }
    ]

    for s in slides:
        p_sec = doc.add_paragraph()
        r_sec = p_sec.add_run(f"📌 {s['slide_no']}: {s['section_title']}")
        r_sec.font.name = "Arial"
        r_sec.font.size = Pt(13)
        r_sec.font.bold = True
        r_sec.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        p_sec.paragraph_format.space_before = Pt(14)
        p_sec.paragraph_format.space_after = Pt(4)

        table = doc.add_table(rows=4, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        rows_def = [
            ("Slide Topic / หัวข้อสไลด์", s['topic']),
            ("English Speech Script\n(บทพูดภาษาอังกฤษ)", s['speech']),
            ("Thai Translation\n(คำแปลและความหมายภาษาไทย)", s['translation']),
            ("Speaking & Vocal Tips\n(คำแนะนำการพูดพรีเซ้นต์)", s['tips'])
        ]

        for r_i, (lbl, val) in enumerate(rows_def):
            c_lbl = table.cell(r_i, 0)
            c_val = table.cell(r_i, 1)

            c_lbl.width = Inches(2.0)
            c_val.width = Inches(4.6)

            set_cell_background(c_lbl, "1E3A8A" if r_i == 0 else "F1F5F9")
            set_cell_background(c_val, "EFF6FF" if r_i == 1 else ("FFFFFF" if r_i % 2 == 0 else "F8FAFC"))

            set_cell_margins(c_lbl, top=100, bottom=100, left=120, right=120)
            set_cell_margins(c_val, top=100, bottom=100, left=120, right=120)

            p_l = c_lbl.paragraphs[0]
            r_l = p_l.add_run(lbl)
            r_l.font.name = "Arial"
            r_l.font.size = Pt(9.5)
            r_l.font.bold = True
            r_l.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r_i == 0 else RGBColor(0x1E, 0x3A, 0x8A)

            p_v = c_val.paragraphs[0]
            r_v = p_v.add_run(val)
            r_v.font.name = "Arial"
            r_v.font.size = Pt(10)
            if r_i == 1:
                r_v.font.bold = True
                r_v.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif r_i == 3:
                r_v.font.italic = True
                r_v.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
            else:
                r_v.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Vocabulary Table Section
    p_vhdr = doc.add_paragraph()
    r_vh = p_vhdr.add_run("💡 KEY TECHNICAL VOCABULARY & PRONUNCIATION (ตารางคำศัพท์และการออกเสียง)")
    r_vh.font.name = "Arial"
    r_vh.font.size = Pt(13)
    r_vh.font.bold = True
    r_vh.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p_vhdr.paragraph_format.space_before = Pt(14)
    p_vhdr.paragraph_format.space_after = Pt(6)

    vt = doc.add_table(rows=9, cols=3)
    vt.alignment = WD_TABLE_ALIGNMENT.CENTER
    vt.autofit = False

    vh = ["English Term", "Thai Translation", "Pronunciation / Usage Guide"]
    for c_i, h in enumerate(vh):
        c = vt.cell(0, c_i)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    v_items = [
        ("Smart Job Matching Platform", "แพลตฟอร์มคัดกรองจัดหางานอัจฉริยะ", "สมาร์ต-จ็อบ-แมตชิง-แพลตฟอร์ม"),
        ("Fresh Graduates", "นักศึกษาจบใหม่ / ผู้สำเร็จการศึกษาใหม่", "เฟรช-แกรด-จู-เอทส์"),
        ("Mismatch", "ความไม่สอดคล้อง / ไม่ตรงกันของทักษะ", "มิส-แมตช์"),
        ("Single Page Application (SPA)", "เว็บแอปพลิเคชันหน้าเดียวที่โหลดรวดเร็ว", "ซิงเกิล-เพจ-แอปพลิเคชัน"),
        ("RESTful API", "สถาปัตยกรรมเชื่อมต่อข้อมูลหน้าเว็บกับเซิร์ฟเวอร์", "เรสต์ฟูล-เอพีไอ"),
        ("Role-Based Access Control (RBAC)", "ระบบควบคุมสิทธิ์ตามบทบาท (User/Employer/Admin)", "โรล-เบสด์-แอคเซส-คอนโทรล"),
        ("Dynamic AI Matching Engine", "ระบบเอไอคำนวณอัตราความเข้ากันได้แบบเรียลไทม์", "ไดนามิก-เอไอ-แมตชิง-เอนจิน"),
        ("Admin Moderation System", "ระบบผู้ดูแลระบบสำหรับกำกับดูแลความปลอดภัย", "แอดมิน-มอดเดอเรชัน-ซิสเต็ม")
    ]

    for r_i, (term, trans, pron) in enumerate(v_items, start=1):
        bg = "F8FAFC" if r_i % 2 == 1 else "FFFFFF"
        for c_i, text_val in enumerate([term, trans, pron]):
            c = vt.cell(r_i, c_i)
            c.width = [Inches(2.4), Inches(2.4), Inches(1.8)][c_i]
            set_cell_background(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            r = p.add_run(text_val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            if c_i == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            elif c_i == 2:
                r.font.italic = True
                r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    try:
        doc.save(output_path)
        print("SUCCESS: Enhanced Word document saved at " + output_path)
    except PermissionError:
        alt_path = output_path.replace(".docx", "_Updated.docx")
        doc.save(alt_path)
        print("SUCCESS: Enhanced Word document saved at " + alt_path)

if __name__ == "__main__":
    create_enhanced_docx("d:/work90/bluehouse/BlueHouse_Jobs_Presentation_Script.docx")
