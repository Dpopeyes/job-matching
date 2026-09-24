import json
import sqlite3
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# ----------------------------------------------------
# Sheet 1: AI Fine-Tuning Dataset
# ----------------------------------------------------
ws1 = wb.create_sheet(title="AI Fine-Tuning Dataset")

headers1 = [
    "ลำดับ (No.)",
    "ข้อมูลผู้สมัครและทักษะ (Input Text)",
    "สายงานแนะนำหลัก (Suggested Career)",
    "สายงานสำรอง (Alternative Career)",
    "ตำแหน่งงานแนะนำ (Recommended Jobs)",
    "เหตุผลการวิเคราะห์ของ AI (Match Reason)"
]

ws1.append(headers1)

jsonl_path = "d:/work90/bluehouse/gemini_fine_tune_dataset.jsonl"
if os.path.exists(jsonl_path):
    with open(jsonl_path, "r", encoding="utf-8") as f:
        idx = 1
        for line in f:
            if not line.trim() if hasattr(line, 'trim') else not line.strip():
                continue
            item = json.loads(line)
            inp = item.get("input_text", "")
            out_str = item.get("output_text", "{}")
            try:
                out = json.loads(out_str)
            except:
                out = {}
            
            rec_jobs = ", ".join(out.get("recommendedJobTitles", []))
            
            ws1.append([
                idx,
                inp,
                out.get("suggestedCareerPath", ""),
                out.get("alternativeCareerPath", ""),
                rec_jobs,
                out.get("careerMatchReason", "")
            ])
            idx += 1

# ----------------------------------------------------
# Sheet 2: Jobs Database ( SQLite Data )
# ----------------------------------------------------
ws2 = wb.create_sheet(title="Database - Jobs")

headers2 = [
    "รหัสงาน (ID)",
    "ชื่อตำแหน่งงาน (Title)",
    "ชื่อบริษัท/องค์กร (Company)",
    "หมวดหมู่งาน (Category)",
    "สถานที่ทำงาน (Location)",
    "ประเภทงาน (Type)",
    "อัตราเงินเดือน (Salary)",
    "ทักษะที่ต้องการ (Skills Required)",
    "สถานะการอนุมัติ (Approval Status)",
    "ประเภทนายจ้าง (Employer Type)"
]
ws2.append(headers2)

db_path = "d:/work90/bluehouse/server/database.sqlite"
if os.path.exists(db_path):
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, company, category, location, type, salary, skillsRequired, approvalStatus, employerType FROM jobs ORDER BY rowid DESC")
    rows = cursor.fetchall()
    for r in rows:
        skills = r[7]
        try:
            skills_list = json.loads(skills)
            skills_fmt = ", ".join(skills_list)
        except:
            skills_fmt = str(skills)
        ws2.append([r[0], r[1], r[2], r[3], r[4], r[5], r[6], skills_fmt, r[8], r[9]])

# ----------------------------------------------------
# Sheet 3: Users Database ( SQLite Data )
# ----------------------------------------------------
ws3 = wb.create_sheet(title="Database - Users")

headers3 = [
    "รหัสผู้ใช้ (ID)",
    "ชื่อ-นามสกุล (Name)",
    "อีเมล (Email)",
    "บทบาท (Role)",
    "มหาวิทยาลัย/องค์กร",
    "สาขาวิชา/อุตสาหกรรม",
    "ประเภทนายจ้าง (Employer Type)",
    "สถานะยืนยันตัวตน (KYC Status)"
]
ws3.append(headers3)

if os.path.exists(db_path):
    cursor.execute("SELECT id, name, email, role, university, major, employerType, faceKYCVerified FROM users ORDER BY rowid DESC")
    u_rows = cursor.fetchall()
    for u in u_rows:
        kyc = "ยืนยันเรียบร้อย (Verified)" if u[7] == 1 else "ยังไม่อยู่ในระบบ"
        ws3.append([u[0], u[1], u[2], u[3], u[4], u[5], u[6], kyc])
    conn.close()


# ----------------------------------------------------
# Styling Excel Sheets
# ----------------------------------------------------
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
body_font = Font(name="Tahoma", size=10)
border_side = Side(border_style="thin", color="CBD5E1")
thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

for sheet in wb.worksheets:
    sheet.views.sheetView[0].showGridLines = True
    
    # Format Headers
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Format Rows
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            
    # Auto Column Widths
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        sheet.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 50)

output_excel = "d:/work90/bluehouse/BlueHouse_Dataset_and_Database.xlsx"
wb.save(output_excel)
print(f"SUCCESS: Exported Excel file to {output_excel}")
